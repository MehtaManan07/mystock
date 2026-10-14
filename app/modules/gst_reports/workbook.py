"""GSTR-1 Excel interchange layout, to be validated by the official offline tool."""
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, ROUND_FLOOR
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from app.core.exceptions import ValidationError
from .reporting import STATES, UQC_LABELS, series_rows
from .schemas import Report


# Header names/order follow the supplied TallyPrime May-2025+ export.
HEADERS = {
    "b2b,sez,de": ["GSTIN/UIN of Recipient", "Receiver Name", "Invoice Number", "Invoice date", "Invoice Value", "Place Of Supply", "Reverse Charge", "Applicable % of Tax Rate", "Invoice Type", "E-Commerce GSTIN", "Rate", "Taxable Value", "Cess Amount"],
    "b2ba": ["GSTIN/UIN of Recipient", "Receiver Name", "Original Invoice Number", "Original Invoice date", "Revised Invoice Number", "Revised Invoice date", "Invoice Value", "Place Of Supply", "Reverse Charge", "Applicable % of Tax Rate", "Invoice Type", "E-Commerce GSTIN", "Rate", "Taxable Value", "Cess Amount"],
    "b2cl": ["Invoice Number", "Invoice date", "Invoice Value", "Place Of Supply", "Applicable % of Tax Rate", "Rate", "Taxable Value", "Cess Amount", "E-Commerce GSTIN"],
    "b2cla": ["Original Invoice Number", "Original Invoice date", "Original Place Of Supply", "Revised Invoice Number", "Revised Invoice date", "Invoice Value", "Applicable % of Tax Rate", "Rate", "Taxable Value", "Cess Amount", "E-Commerce GSTIN"],
    "b2cs": ["Type", "Place Of Supply", "Applicable % of Tax Rate", "Rate", "Taxable Value", "Cess Amount", "E-Commerce GSTIN"],
    "b2csa": ["Financial Year", "Original Month", "Place Of Supply", "Type", "Applicable % of Tax Rate", "Rate", "Taxable Value", "Cess Amount", "E-Commerce GSTIN"],
    "cdnr": ["GSTIN/UIN of Recipient", "Receiver Name", "Note Number", "Note Date", "Note Type", "Place Of Supply", "Reverse Charge", "Note Supply Type", "Note Value", "Applicable % of Tax Rate", "Rate", "Taxable Value", "Cess Amount"],
    "cdnra": ["GSTIN/UIN of Recipient", "Receiver Name", "Original Note Number", "Original Note Date", "Revised Note Number", "Revised Note Date", "Note Type", "Place Of Supply", "Reverse Charge", "Note Supply Type", "Note Value", "Applicable % of Tax Rate", "Rate", "Taxable Value", "Cess Amount"],
    "cdnur": ["UR Type", "Note Number", "Note Date", "Note Type", "Place Of Supply", "Note Value", "Applicable % of Tax Rate", "Rate", "Taxable Value", "Cess Amount"],
    "cdnura": ["UR Type", "Original Note Number", "Original Note Date", "Revised Note Number", "Revised Note Date", "Note Type", "Place Of Supply", "Note Value", "Applicable % of Tax Rate", "Rate", "Taxable Value", "Cess Amount"],
    "exp": ["Export Type", "Invoice Number", "Invoice date", "Invoice Value", "Port Code", "Shipping Bill Number", "Shipping Bill Date", "Rate", "Taxable Value", "Cess Amount"],
    "expa": ["Export Type", "Original Invoice Number", "Original Invoice date", "Revised Invoice Number", "Revised Invoice date", "Invoice Value", "Port Code", "Shipping Bill Number", "Shipping Bill Date", "Rate", "Taxable Value", "Cess Amount"],
    "at": ["Place Of Supply", "Applicable % of Tax Rate", "Rate", "Gross Advance Received", "Cess Amount"],
    "ata": ["Financial Year", "Original Month", "Original Place Of Supply", "Applicable % of Tax Rate", "Rate", "Gross Advance Received", "Cess Amount"],
    "atadj": ["Place Of Supply", "Applicable % of Tax Rate", "Rate", "Gross Advance Adjusted", "Cess Amount"],
    "atadja": ["Financial Year", "Original Month", "Original Place Of Supply", "Applicable % of Tax Rate", "Rate", "Gross Advance Adjusted", "Cess Amount"],
    "exemp": ["Description", "Nil Rated Supplies", "Exempted (other than nil rated/non GST supply )", "Non-GST supplies"],
    "hsn(b2b)": ["HSN", "Description", "UQC", "Total Quantity", "Total Value", "Rate", "Taxable Value", "Integrated Tax Amount", "Central Tax Amount", "State/UT Tax Amount", "Cess Amount"],
    "hsn(b2c)": ["HSN", "Description", "UQC", "Total Quantity", "Total Value", "Rate", "Taxable Value", "Integrated Tax Amount", "Central Tax Amount", "State/UT Tax Amount", "Cess Amount"],
    "docs": ["Nature of Document", "Sr. No. From", "Sr. No. To", "Total Number", "Cancelled"],
}


def allocate(target: Decimal, weights: list[Decimal]) -> list[Decimal]:
    """Allocate a billed tax component in paise, preserving its exact total."""
    total = sum(weights, Decimal(0))
    if total == 0:
        if target:
            raise ValidationError("Tax cannot be allocated to zero-valued HSN tax groups")
        return [Decimal(0) for _ in weights]
    shares = [target * weight / total for weight in weights]
    result = [share.quantize(Decimal(".01"), rounding=ROUND_FLOOR) for share in shares]
    pennies = int((target - sum(result, Decimal(0))) * 100)
    order = sorted(range(len(weights)), key=lambda index: (-(shares[index] - result[index]), index))
    for index in order[:pennies]:
        result[index] += Decimal(".01")
    if sum(result, Decimal(0)) != target:
        raise ValidationError("HSN allocation does not reconcile to the billed tax component")
    return result


def export_rows(report: Report) -> dict[str, list[list]]:
    if not report.exportable:
        raise ValidationError("GST export blocked. Resolve every reconciliation blocker first.")
    rows = {name: [] for name in HEADERS}
    b2cs = defaultdict(Decimal)
    exempt = defaultdict(lambda: [Decimal(0), Decimal(0), Decimal(0)])
    hsn = defaultdict(lambda: [Decimal(0) for _ in range(6)])
    taxable_written = Decimal(0)
    for view in report.documents:
        if view.disposition != "include":
            continue
        doc = view.document
        pos = f"{doc.place_of_supply}-{STATES[doc.place_of_supply]}"
        interstate = doc.place_of_supply != doc.seller_gstin[:2]
        registered = doc.customer_gstin is not None
        sign = -1 if doc.kind == "credit_note" else 1
        note_type = "C" if doc.kind == "credit_note" else "D"
        dated = doc.date
        taxable_by_rate = defaultdict(Decimal)
        hsn_groups = defaultdict(lambda: {"qty": Decimal(0), "value": Decimal(0), "tax": Decimal(0)})
        for line in doc.lines:
            if line.uqc not in UQC_LABELS:
                raise ValidationError(f"{view.source_key}: UQC {line.uqc} is not mapped by this Excel layout version.")
            if line.treatment == "taxable":
                taxable_by_rate[line.tax_rate] += line.taxable_value
            else:
                label = f"{'Inter' if interstate else 'Intra'}-State supplies to {'registered' if registered else 'unregistered'} persons"
                column = {"nil_rated": 0, "exempt": 1, "non_gst": 2}[line.treatment]
                exempt[label][column] += sign * line.taxable_value
            if line.treatment != "non_gst":
                group = hsn_groups[(line.hsn, line.uqc, line.tax_rate)]
                group["qty"] += line.quantity
                group["value"] += line.taxable_value
                group["tax"] += line.tax_amount
        for rate, value in sorted(taxable_by_rate.items()):
            taxable_written += sign * value
            if doc.kind == "invoice" and registered:
                rows["b2b,sez,de"].append([
                    doc.customer_gstin, doc.customer_name, doc.number, dated, doc.total_amount,
                    pos, "Y" if doc.reverse_charge else "N", None, "Regular B2B", None, rate, value, 0,
                ])
            elif doc.kind == "invoice" and interstate and doc.total_amount > 100000:
                rows["b2cl"].append([doc.number, dated, doc.total_amount, pos, None, rate, value, 0, None])
            elif doc.kind != "invoice" and registered:
                rows["cdnr"].append([
                    doc.customer_gstin, doc.customer_name, doc.number, dated, note_type, pos,
                    "Y" if doc.reverse_charge else "N", "Regular B2B", doc.total_amount, None, rate, value, 0,
                ])
            elif doc.kind != "invoice" and interstate and doc.original_invoice_value > 100000:
                rows["cdnur"].append(["B2CL", doc.number, dated, note_type, pos, doc.total_amount, None, rate, value, 0])
            else:
                b2cs[(pos, rate)] += sign * value
        for bucket in doc.rate_taxes:
            keys = sorted(key for key in hsn_groups if key[2] == bucket.rate)
            weights = [hsn_groups[key]["tax"] for key in keys]
            central = allocate(bucket.cgst, weights)
            for index, key in enumerate(keys):
                group = hsn_groups[key]
                target = hsn[("hsn(b2b)" if registered else "hsn(b2c)", *key)]
                igst = group["tax"] if doc.tax_type == "igst" else Decimal(0)
                cgst = central[index]
                sgst = group["tax"] - cgst if doc.tax_type == "cgst_sgst" else Decimal(0)
                values = (group["qty"], group["value"], igst, cgst, sgst, group["value"] + group["tax"])
                for column, value in enumerate(values):
                    target[column] += sign * value
    for (pos, rate), value in sorted(b2cs.items()):
        if value < 0:
            raise ValidationError("B2CS net value is negative; obtain period-specific filing guidance before exporting.")
        rows["b2cs"].append(["OE", pos, None, rate, value, 0, None])
    for interstate in (True, False):
        for registered in (True, False):
            label = f"{'Inter' if interstate else 'Intra'}-State supplies to {'registered' if registered else 'unregistered'} persons"
            if any(value < 0 for value in exempt[label]):
                raise ValidationError("Negative exemption-summary values require period-specific filing guidance.")
            rows["exemp"].append([label, *exempt[label]])
    for (sheet, code, uqc, rate), values in sorted(hsn.items()):
        qty, value, igst, cgst, sgst, total = values
        if any(amount < 0 for amount in values):
            raise ValidationError("Negative net HSN quantities/amounts require period-specific filing guidance.")
        rows[sheet].append([code, None, UQC_LABELS[uqc], qty, total, rate, value, igst, cgst, sgst, 0])
    rows["docs"], problems = series_rows(report.documents)
    if problems:
        raise ValidationError("Document series reconciliation failed")
    if taxable_written != report.totals.taxable:
        raise ValidationError("Export taxable turnover differs from the reconciliation report")
    for column, field in ((7, "igst"), (8, "cgst"), (9, "sgst")):
        actual = sum((row[column] for name in ("hsn(b2b)", "hsn(b2c)") for row in rows[name]), Decimal(0))
        if actual != getattr(report.totals, field):
            raise ValidationError(f"HSN {field} does not reconcile to the billed tax")
    hsn_value = sum((row[6] for name in ("hsn(b2b)", "hsn(b2c)") for row in rows[name]), Decimal(0))
    if hsn_value != report.totals.taxable + report.totals.nil_rated + report.totals.exempt:
        raise ValidationError("HSN turnover does not reconcile to classified supplies")
    for index, field in ((1, "nil_rated"), (2, "exempt"), (3, "non_gst")):
        if sum((row[index] for row in rows["exemp"]), Decimal(0)) != getattr(report.totals, field):
            raise ValidationError(f"Exemption summary {field} differs from the report")
    return rows


def generate_workbook(report: Report) -> bytes:
    rows = export_rows(report)
    workbook = Workbook()
    help_sheet = workbook.active
    help_sheet.title = "Help Instruction"
    instructions = [
        "GSTR-1 offline-tool interchange workbook, not a filed return.",
        f"GSTIN: {report.seller_gstin}; period: {report.request.from_date} to {report.request.to_date}.",
        f"Reconciliation SHA256: {report.report_sha256}",
        "Validate/import with the current official GST Returns Offline Tool and review its errors and generated JSON.",
        "Do not upload this Excel directly to the GST portal. Upload the official tool's JSON after review.",
        "Invoice value repeats for each tax-rate row; count/sum unique invoices, not B2B rows.",
        "HSN Total Value excludes invoice round-off; the reconciliation report retains it separately.",
        "Non-GST values are in exemp, not the HSN summary.",
        "Verify exemp and docs in the portal separately if the offline tool does not import them.",
        "No IFF, previously filed invoices, amendments, ECO, exports/SEZ, advances or cess are included.",
    ]
    for index, value in enumerate(instructions, 1):
        help_sheet.cell(index, 1, value)
    help_sheet.column_dimensions["A"].width = 125
    for name, headers in HEADERS.items():
        sheet = workbook.create_sheet(name)
        sheet.cell(1, 1, f"Summary For {name}")
        for column, header in enumerate(headers, 1):
            cell = sheet.cell(4, column, header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9E1F2")
        for index, row in enumerate(rows[name], 5):
            for column, value in enumerate(row, 1):
                cell = sheet.cell(index, column, value)
                if isinstance(value, str):
                    # Prevent customer names/descriptions from becoming formulas.
                    cell.data_type = "s"
                if isinstance(value, Decimal):
                    cell.number_format = "0.000" if headers[column - 1] == "Total Quantity" else "0.00"
                elif isinstance(value, date):
                    cell.number_format = "dd-mmm-yyyy"
        sheet.freeze_panes = "A5"
    content = BytesIO()
    workbook.save(content)
    workbook.close()
    data = content.getvalue()
    # Round-trip the actual output, including text identifiers and numeric cells.
    reopened = load_workbook(BytesIO(data), read_only=True, data_only=False)
    try:
        if reopened.sheetnames != ["Help Instruction", *HEADERS]:
            raise ValidationError("Generated GST worksheet list differs from the supported layout")
        for name, headers in HEADERS.items():
            sheet = reopened[name]
            if [cell.value for cell in sheet[4]] != headers:
                raise ValidationError(f"Generated GST headers differ in {name}")
            if not rows[name]:
                continue
            actual_rows = sheet.iter_rows(min_row=5, max_row=4 + len(rows[name]), max_col=len(headers))
            for actual_row, row in zip(actual_rows, rows[name], strict=True):
                for cell, expected in zip(actual_row, row, strict=True):
                    actual = cell.value
                    if isinstance(expected, Decimal):
                        equal = actual is not None and Decimal(str(actual)) == expected
                    elif isinstance(expected, date):
                        equal = (actual.date() if isinstance(actual, datetime) else actual) == expected
                    else:
                        equal = actual == expected
                    if not equal or cell.data_type == "f":
                        raise ValidationError(f"GST output round-trip mismatch at {name}!{cell.coordinate}")
    finally:
        reopened.close()
    return data
