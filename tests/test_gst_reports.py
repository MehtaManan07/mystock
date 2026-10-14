"""Offline GST reporting regressions. Run only with user permission."""
from datetime import date
from decimal import Decimal
from io import BytesIO
import os
import unittest
from unittest.mock import AsyncMock, patch

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError as SchemaError
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

with patch.dict(os.environ, {
    "TURSO_DATABASE_URL": "libsql://unit-test.invalid",
    "TURSO_AUTH_TOKEN": "not-a-real-token",
    "GCP_BUCKET_NAME": "unit-tests",
    "SECRET_KEY": "unit-test-only",
}):
    from app.main import app as application
    from app.core.db.base import Base
    from app.core.db.engine import engine as configured_engine
    from app.core.exceptions import ConflictError, ValidationError
    from app.core.response_interceptor import SuccessResponseInterceptor
    from app.modules.contacts.models import Contact
    from app.modules.gst_reports.models import GstReview
    from app.modules.gst_reports.reporting import ATTESTATIONS, build_report, valid_gstin
    from app.modules.gst_reports.router import router
    from app.modules.gst_reports.schemas import (
        AdditionalDocumentInput, GstDocument, GstLine, RateTax, ReportRequest, ReviewInput, SourceView,
    )
    from app.modules.gst_reports.service import GstReportingService
    from app.modules.gst_reports.sources import capture_source, read_source
    from app.modules.gst_reports.workbook import HEADERS, allocate, export_rows, generate_workbook
    from app.modules.pdf_invoices.schemas import PdfInvoiceMetadata
    from app.modules.pdf_invoices.service import PdfInvoiceService
    from app.modules.products.models import Product
    from app.modules.settings.models import CompanySettings
    from app.modules.transactions.models import Transaction, TransactionItem
    from app.modules.transactions.schemas import CreateSaleDto
    from app.modules.transactions.service import TransactionsService
    from app.modules.containers.models import Container
    from app.modules.container_products.models import ContainerProduct
    from app.modules.users.auth import TokenData, get_current_user
    from app.modules.users.models import User


# Synthetic checksum-valid identifiers, not customer registration data.
SELLER = "24AAAAA0000A1Z8"
BUYER = "24BBBBB0000B1ZZ"


def request(**changes):
    values = dict(
        from_date="2025-07-01", to_date="2025-09-30", filing_frequency="quarterly",
        preceding_year_turnover_above_5cr=False, **dict.fromkeys(ATTESTATIONS, True),
    )
    values.update(changes)
    return ReportRequest(**values)


def document(number="INV-001", **changes):
    values = dict(
        number=number, date="2025-09-15", seller_gstin=SELLER,
        customer_name="Example Buyer", customer_gstin=BUYER, place_of_supply="24",
        invoice_type="regular", reverse_charge=False, tax_type="cgst_sgst",
        subtotal="100", tax_amount="18", cgst="9", sgst="9", igst="0",
        total_amount="118", series="INV",
        lines=[GstLine(
            line_key="1", description="Goods", quantity=10, unit_price="10", taxable_value="100",
            tax_rate="18", tax_amount="18", hsn="44111200", uqc="NOS", treatment="taxable",
        )],
        rate_taxes=[RateTax(rate=18, cgst=9, sgst=9)],
    )
    values.update(changes)
    return GstDocument(**values)


def view(doc=None, key="transaction:1", **changes):
    return SourceView(
        source_key=key, source_sha256="a" * 64, revision=1, verified=True,
        source_deleted=changes.pop("source_deleted", False), upload_ready=True,
        stale_review=changes.pop("stale_review", False),
        disposition=changes.pop("disposition", "include"), document=doc or document(), **changes,
    )


class CalculationTests(unittest.TestCase):
    def test_valid_gstin_and_period_validation(self):
        self.assertTrue(valid_gstin(SELLER))
        self.assertTrue(valid_gstin(BUYER))
        self.assertFalse(valid_gstin(BUYER[:-1] + "0"))
        for values in (
            dict(from_date="2025-07-02"), dict(to_date="2025-09-29"),
            dict(from_date="2025-08-01", to_date="2025-10-31"),
            dict(filing_frequency="monthly"),
        ):
            with self.subTest(values=values), self.assertRaises(SchemaError):
                request(**values)
        self.assertEqual(request(filing_frequency="monthly", to_date="2025-07-31").to_date, date(2025, 7, 31))

    def test_mixed_rates_repeat_invoice_value_and_reconcile_hsn(self):
        second = GstLine(line_key="2", description="Other", quantity=1, taxable_value="50",
                         tax_rate="5", tax_amount="2.50", hsn="39264099", uqc="SET", treatment="taxable")
        doc = document(subtotal="150", tax_amount="20.50", total_amount="171", round_off=".50",
                       cgst="10.25", sgst="10.25",
                       lines=[*document().lines, second],
                       rate_taxes=[RateTax(rate=18, cgst=9, sgst=9), RateTax(rate=5, cgst="1.25", sgst="1.25")])
        report = build_report(request(), SELLER, [view(doc)])
        self.assertTrue(report.exportable, report.issues)
        rows = export_rows(report)
        self.assertEqual(len(rows["b2b,sez,de"]), 2)
        self.assertEqual([row[4] for row in rows["b2b,sez,de"]], [Decimal("171")] * 2)
        self.assertEqual(sum(row[11] for row in rows["b2b,sez,de"]), Decimal("150"))
        self.assertEqual(sum(row[4] for row in rows["hsn(b2b)"]), Decimal("170.50"))
        self.assertEqual(report.totals.round_off, Decimal(".50"))

    def test_b2c_threshold_is_one_lakh_and_not_old_two_lakh_fifty(self):
        docs = []
        for index, value in enumerate((100000, 100001), 1):
            taxable = Decimal(value) / Decimal("1.18")
            taxable = taxable.quantize(Decimal(".01"))
            tax = (taxable * Decimal(".18")).quantize(Decimal(".01"))
            doc = document(
                number=f"INV-{index:03d}", customer_gstin=None, place_of_supply="37", tax_type="igst",
                subtotal=taxable, tax_amount=tax, igst=tax, cgst=0, sgst=0,
                total_amount=value, round_off=Decimal(value) - taxable - tax,
                lines=[GstLine(line_key="1", description="Goods", quantity=1, taxable_value=taxable,
                               tax_rate=18, tax_amount=tax, hsn="44111200", uqc="NOS", treatment="taxable")],
                rate_taxes=[RateTax(rate=18, igst=tax)],
            )
            docs.append(view(doc, key=f"transaction:{index}"))
        report = build_report(request(), SELLER, docs)
        self.assertTrue(report.exportable, report.issues)
        rows = export_rows(report)
        self.assertEqual(len(rows["b2cs"]), 1)
        self.assertEqual(rows["b2cl"][0][0], "INV-002")
        self.assertEqual(len(rows["hsn(b2c)"]), 1)

    def test_credit_debit_and_exempt_non_gst_netting(self):
        credit = document(number="CN-001", kind="credit_note", series="CN",
                          original_invoice_number="INV-001", original_invoice_date="2025-09-15")
        debit = document(number="DN-001", kind="debit_note", series="DN",
                         original_invoice_number="INV-001", original_invoice_date="2025-09-15")
        exempt_lines = [
            GstLine(line_key=str(index), description=kind, quantity=1, taxable_value="10",
                    tax_rate=0, tax_amount=0, hsn="44111200", uqc="NOS", treatment=kind)
            for index, kind in enumerate(("exempt", "non_gst", "nil_rated"), 1)
        ]
        exempt = document(number="EX-001", series="EX", subtotal=30, tax_amount=0, cgst=0, sgst=0,
                          total_amount=30, lines=exempt_lines, rate_taxes=[RateTax(rate=0)])
        docs = [view(), view(credit, key="additional:cn"), view(debit, key="additional:dn"), view(exempt, key="additional:ex")]
        report = build_report(request(), SELLER, docs)
        self.assertTrue(report.exportable, report.issues)
        self.assertEqual(report.totals.total, Decimal("148"))
        self.assertEqual((report.totals.exempt, report.totals.non_gst, report.totals.nil_rated), (Decimal(10),) * 3)
        rows = export_rows(report)
        self.assertEqual({row[4] for row in rows["cdnr"]}, {"C", "D"})
        self.assertEqual(len(rows["docs"]), 4)
        self.assertEqual(sum(row[6] for row in rows["hsn(b2b)"]), Decimal("120"))

    def test_fail_closed_on_missing_details_or_false_attestation(self):
        for mutation, code in (
            ({"hsn": None}, "hsn"), ({"uqc": None}, "uqc"),
            ({"tax_rate": None}, "missing_tax_snapshot"),
            ({"treatment": "unknown"}, "supply_classification"),
            ({"tax_amount": Decimal("17")}, "line_tax_arithmetic"),
        ):
            doc = document()
            doc.lines[0] = doc.lines[0].model_copy(update=mutation)
            report = build_report(request(), SELLER, [view(doc)])
            self.assertIn(code, {issue.code for issue in report.issues})
            with self.assertRaises(ValidationError):
                generate_workbook(report)
        report = build_report(request(sources_complete=False), SELLER, [view()])
        self.assertFalse(report.exportable)
        self.assertIn("sources_complete", {issue.code for issue in report.issues})

    def test_duplicate_identity_stale_review_and_series_gaps_block(self):
        docs = [view(), view(document(), key="pdf_invoice:1")]
        self.assertIn("duplicate_identity", {i.code for i in build_report(request(), SELLER, docs).issues})
        report = build_report(request(), SELLER, [view(stale_review=True)])
        self.assertIn("stale_review", {i.code for i in report.issues})
        docs = [view(), view(document("INV-003"), key="transaction:3")]
        self.assertIn("series_gap", {i.code for i in build_report(request(), SELLER, docs).issues})
        docs.append(view(document("INV-002"), key="transaction:2", source_deleted=True, disposition="cancelled"))
        report = build_report(request(), SELLER, docs)
        self.assertTrue(report.exportable, report.issues)
        self.assertEqual(export_rows(report)["docs"], [["Invoices for outward supply", "INV-001", "INV-003", 3, 1]])
        self.assertEqual(report.source_invoice_total, report.included_document_total + report.excluded_document_total)

    def test_decimal_allocation_and_workbook_roundtrip_no_formula(self):
        self.assertEqual(sum(allocate(Decimal(".03"), [Decimal(".01")] * 5)), Decimal(".03"))
        doc = document(customer_name="=HYPERLINK(\"https://example.invalid\")")
        report = build_report(request(), SELLER, [view(doc)])
        content = generate_workbook(report)
        book = load_workbook(BytesIO(content), data_only=False)
        self.assertEqual(book.sheetnames, ["Help Instruction", *HEADERS])
        self.assertEqual(book["b2b,sez,de"]["B5"].data_type, "s")
        self.assertEqual(book["b2b,sez,de"]["A5"].value, BUYER)
        self.assertEqual(book["b2b,sez,de"]["L5"].value, 100)
        self.assertEqual(book["b2b,sez,de"]["D5"].value.date(), doc.date)
        self.assertEqual(book["b2b,sez,de"]["D5"].number_format, "dd-mmm-yyyy")
        book.close()

    def test_header_component_mismatch_and_unsupported_uqc_block(self):
        doc = document(cgst="8.99", sgst="9.01")
        report = build_report(request(), SELLER, [view(doc)])
        self.assertIn("component_reconciliation", {issue.code for issue in report.issues})
        doc = document()
        doc.lines[0].uqc = "BAG"
        report = build_report(request(), SELLER, [view(doc)])
        self.assertIn("unsupported_uqc", {issue.code for issue in report.issues})

    def test_negative_net_hsn_is_a_reconciliation_blocker(self):
        note = document("CN-001", kind="credit_note", series="CN",
                        original_invoice_number="PRIOR-001", original_invoice_date="2025-06-30")
        report = build_report(request(), SELLER, [view(note, key="additional:cn")])
        self.assertFalse(report.exportable)
        self.assertIn("export_layout", {issue.code for issue in report.issues})

    def test_cumulative_credit_cannot_exceed_original_and_debit_notes(self):
        notes = [document(f"CN-{index:03d}", kind="credit_note", series="CN",
                          original_invoice_number="INV-001", original_invoice_date="2025-09-15")
                 for index in (1, 2)]
        report = build_report(request(), SELLER, [
            view(), *[view(note, key=f"additional:{index}") for index, note in enumerate(notes)],
        ])
        self.assertIn("excess_credit", {issue.code for issue in report.issues})


class DatabaseTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        guard = patch.object(configured_engine, "connect", side_effect=AssertionError("Production DB forbidden"))
        guard.start()
        self.addCleanup(guard.stop)
        self.engine = create_engine("sqlite://")
        self.addCleanup(self.engine.dispose)
        Base.metadata.create_all(self.engine)
        with Session(self.engine) as db:
            db.add(User(id=1, username="reviewer", password="unused", name="Reviewer", role="ADMIN"))
            db.add(CompanySettings(id=1, seller_gstin=SELLER))
            db.add(Contact(id=1, name="Example Buyer", phone="", type="customer", gstin=BUYER, balance=0))
            db.add(Product(id=1, name="Goods", display_name="Goods", size="NA", packing="1", hsn_code="44111200", gst_rate=18))
            db.add(Transaction(id=1, transaction_number="INV-001", transaction_date=date(2025, 9, 15),
                               type="sale", contact_id=1, subtotal=100, tax_amount=18, total_amount=118,
                               tax_type="cgst_sgst", discount_amount=0))
            db.add(TransactionItem(id=1, transaction_id=1, product_id=1, quantity=10, unit_price=10,
                                   line_total=100, tax_rate=18, tax_amount=18))
            db.commit()

        async def local_db(fn):
            with Session(self.engine, expire_on_commit=False) as db:
                result = fn(db)
                db.commit()
                return result
        for module in ("gst_reports.service", "transactions.service", "pdf_invoices.service"):
            patched = patch(f"app.modules.{module}.run_db", side_effect=local_db)
            patched.start()
            self.addCleanup(patched.stop)

    async def test_historical_hsn_not_read_from_current_master_and_reviews_preserve_accounting(self):
        source = await GstReportingService.source("transaction", 1)
        self.assertIsNone(source.document.lines[0].hsn)
        self.assertFalse(source.verified)
        doc = document()
        review = ReviewInput(expected_revision=0, expected_source_sha256=source.source_sha256,
                             document=doc, disposition="include", evidence="Checked against original issued PDF.")
        first = await GstReportingService.review("transaction", 1, review, 1)
        retry = await GstReportingService.review("transaction", 1, review, 1)
        self.assertEqual(first["id"], retry["id"])
        report = await GstReportingService.report(request())
        self.assertTrue(report.exportable, report.issues)
        with Session(self.engine) as db:
            self.assertEqual(db.scalar(select(Contact.balance).where(Contact.id == 1)), 0)
            self.assertEqual(db.scalar(select(Transaction.total_amount).where(Transaction.id == 1)), 118)
            self.assertEqual(len(db.scalars(select(GstReview)).all()), 1)
        changed = doc.model_copy(update={"total_amount": Decimal("119")})
        with self.assertRaises(ValidationError):
            await GstReportingService.review("transaction", 1, review.model_copy(update={"document": changed}), 1)

    async def test_snapshot_survives_product_master_change_and_detects_invoice_mutation(self):
        with Session(self.engine) as db:
            capture_source(db, "transaction", 1)
            db.commit()
            db.get(Product, 1).hsn_code = "39264099"
            db.commit()
        source = await GstReportingService.source("transaction", 1)
        self.assertEqual(source.document.lines[0].hsn, "44111200")
        self.assertFalse(source.stale_review)
        with Session(self.engine) as db:
            db.execute(Transaction.__table__.update().where(Transaction.id == 1).values(total_amount=119))
            db.commit()
        source = await GstReportingService.source("transaction", 1)
        self.assertTrue(source.stale_review)
        report = await GstReportingService.report(request())
        self.assertEqual(report.source_reconciliation_difference, Decimal(1))
        self.assertIn("source_totals", {issue.code for issue in report.issues})

    async def test_invoice_moved_out_of_fiscal_year_remains_a_blocker(self):
        with Session(self.engine) as db:
            capture_source(db, "transaction", 1)
            db.commit()
            db.execute(Transaction.__table__.update().where(Transaction.id == 1).values(transaction_date=date(2026, 9, 15)))
            db.commit()
        report = await GstReportingService.report(request())
        self.assertEqual([source.source_key for source in report.documents], ["transaction:1"])
        self.assertIn("stale_review", {issue.code for issue in report.issues})
        self.assertEqual(report.documents[0].source_document.date, date(2026, 9, 15))

    async def test_invalid_master_hsn_is_captured_but_not_exported(self):
        with Session(self.engine) as db:
            db.get(Product, 1).hsn_code = "TBD"
            db.flush()
            capture_source(db, "transaction", 1)
            db.commit()
        source = await GstReportingService.source("transaction", 1)
        self.assertEqual(source.document.lines[0].hsn, "TBD")
        report = await GstReportingService.report(request())
        self.assertIn("hsn", {issue.code for issue in report.issues})

    async def test_pdf_reservation_captures_once_and_replay_is_idempotent(self):
        with Session(self.engine) as db:
            metadata = PdfInvoiceMetadata(
                invoice_number="PDF-001", invoice_date="2025-09-15", customer_name="Buyer",
                customer_gstin=BUYER, subtotal=100, cgst=9, sgst=9, total_amount=118,
                source="generated", line_items=[dict(name="Goods", quantity=10, unit_price=10, tax_rate=18, hsn_code="44111200")],
            )
            content = b"%PDF-1.4\nfixture\n%%EOF"
            first = PdfInvoiceService._reserve(db, metadata, content, "bill.pdf", 1)
            db.flush()
            second = PdfInvoiceService._reserve(db, metadata, content, "bill.pdf", 1)
            self.assertEqual(first.id, second.id)
            self.assertEqual(len(db.scalars(select(GstReview)).all()), 1)
            source = read_source(db, f"pdf_invoice:{first.id}")
            self.assertEqual(source.document.lines[0].hsn, "44111200")
            self.assertFalse(source.upload_ready)

    async def test_additional_notes_are_revisioned_without_accounting_writes(self):
        note = document("CN-001", kind="credit_note", series="CN",
                        original_invoice_number="INV-001", original_invoice_date="2025-09-15")
        data = AdditionalDocumentInput(key="note-1", expected_revision=0, document=note,
                                       evidence="Supplier-issued credit note checked by accountant.")
        first = await GstReportingService.additional_document(data, 1)
        retry = await GstReportingService.additional_document(data, 1)
        self.assertEqual(first["id"], retry["id"])
        with Session(self.engine) as db:
            self.assertEqual(len(db.execute(select(Transaction.__table__)).all()), 1)
            self.assertEqual(db.scalar(select(Contact.balance).where(Contact.id == 1)), 0)
        with self.assertRaises(ConflictError):
            await GstReportingService.additional_document(data.model_copy(update={"evidence": "A different concurrent review reason."}), 1)
        history = await GstReportingService.history("additional:note-1")
        self.assertEqual(len(history), 1)

    async def test_new_sale_captures_tax_and_hsn_in_the_accounting_transaction(self):
        with Session(self.engine) as db:
            db.add(Container(id=1, name="Stock", type="mixed"))
            db.add(ContainerProduct(container_id=1, product_id=1, quantity=100))
            db.commit()
        sale = await TransactionsService.create_sale(CreateSaleDto(
            transaction_date=date(2025, 9, 16), contact_id=1,
            items=[dict(product_id=1, container_id=1, quantity=10, unit_price=10)],
        ))
        source = await GstReportingService.source("transaction", sale.id)
        self.assertEqual(source.revision, 1)
        self.assertEqual(source.document.lines[0].hsn, "44111200")
        self.assertEqual(source.document.lines[0].tax_rate, Decimal(18))
        self.assertEqual(source.document.cgst, Decimal(9))
        self.assertIsNone(source.document.place_of_supply)
        self.assertFalse(source.verified)
        with Session(self.engine) as db:
            self.assertEqual(db.scalar(select(ContainerProduct.quantity)), 90)
            self.assertEqual(db.scalar(select(Contact.balance).where(Contact.id == 1)), 118)


class RouterTests(unittest.TestCase):
    def test_authentication_roles_download_shape_and_stale_digest(self):
        app = FastAPI()
        app.include_router(router)
        app.add_middleware(SuccessResponseInterceptor)
        report = build_report(request(), SELLER, [view()])
        payload = {"report": request().model_dump(mode="json"), "expected_report_sha256": report.report_sha256}
        with TestClient(app) as client:
            self.assertIn(client.post("/gst-reports/reconcile", json=payload["report"]).status_code, (401, 403))
            app.dependency_overrides[get_current_user] = lambda: TokenData(1, "jobber", "JOBBER")
            self.assertEqual(client.post("/gst-reports/reconcile", json=payload["report"]).status_code, 403)
            app.dependency_overrides[get_current_user] = lambda: TokenData(1, "admin", "ADMIN")
            with patch.object(GstReportingService, "report", new=AsyncMock(return_value=report)):
                response = client.post("/gst-reports/export", json=payload)
                self.assertEqual(response.status_code, 200, response.text if response.status_code != 200 else "")
                self.assertTrue(response.content.startswith(b"PK"))
                self.assertIn("spreadsheetml", response.headers["content-type"])
                self.assertEqual(response.headers["cache-control"], "private, no-store")
                response = client.post("/gst-reports/export", json={**payload, "expected_report_sha256": "0" * 64})
                self.assertEqual(response.status_code, 409)
            blocked = build_report(request(sources_complete=False), SELLER, [view()])
            with patch.object(GstReportingService, "report", new=AsyncMock(return_value=blocked)):
                response = client.post("/gst-reports/export", json={
                    "report": blocked.request.model_dump(mode="json"), "expected_report_sha256": blocked.report_sha256,
                })
                self.assertEqual(response.status_code, 422)


if __name__ == "__main__":
    unittest.main()
